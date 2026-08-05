#!/usr/bin/env python3
"""
Word 文档内容提取与汇总分析
.docx → python-docx
.doc  → antiword / catdoc
.wps  → antiword / catdoc（WPS老格式本质是OLE2）
"""
import subprocess
from pathlib import Path

PDFS_DIR = Path("/home/scada/SmartTwinRes-skills/pdfs")
OUT_DIR = Path("/home/scada/SmartTwinRes-skills/pdf_text_analysis/word_extracted")
OUT_DIR.mkdir(parents=True, exist_ok=True)


def extract_docx(path: Path) -> str:
    """提取 .docx 内容"""
    import docx
    lines = []
    d = docx.Document(str(path))
    for para in d.paragraphs:
        if para.text.strip():
            lines.append(para.text.strip())
    # 表格
    for i, table in enumerate(d.tables, 1):
        lines.append(f"===== 表格 {i} =====")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_doc_with_antiword(path: Path) -> str:
    """用 antiword 提取 .doc 内容"""
    try:
        r = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(path)],
            capture_output=True, text=True, timeout=30
        )
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
        # 退回 catdoc
        r2 = subprocess.run(
            ["catdoc", "-d", "utf-8", str(path)],
            capture_output=True, text=True, timeout=30
        )
        if r2.returncode == 0 and r2.stdout.strip():
            return r2.stdout
        return f"[提取失败: antiword={r.stderr[:100]} catdoc={r2.stderr[:100]}]"
    except Exception as e:
        return f"[提取失败: {e}]"


def main():
    print("=" * 60)
    print("Word 文档提取")
    print("=" * 60)

    doc_files = sorted(PDFS_DIR.rglob("*.doc"))
    docx_files = sorted(PDFS_DIR.rglob("*.docx"))
    wps_files = sorted(PDFS_DIR.rglob("*.wps"))
    all_files = sorted(doc_files + docx_files + wps_files, key=lambda p: str(p))
    print(f"共 {len(all_files)} 个 Word 文档（doc:{len(doc_files)} docx:{len(docx_files)} wps:{len(wps_files)}）\n")

    results = []
    for idx, fp in enumerate(all_files, 1):
        rel = str(fp.relative_to(PDFS_DIR))
        ext = fp.suffix.lower()
        print(f"[{idx}/{len(all_files)}] {rel}", end=" ", flush=True)

        try:
            if ext == ".docx":
                text = extract_docx(fp)
            else:  # .doc / .wps
                text = extract_doc_with_antiword(fp)

            # 判断是否成功（含错误标记则失败）
            if text.startswith("[提取失败"):
                print(f"❌ {text}")
                results.append({"file": rel, "ext": ext, "error": text})
                continue

            # 保存（文件名带路径前缀，避免不同目录同名文件互相覆盖）
            path_prefix = str(fp.relative_to(PDFS_DIR)).replace("/", "_").replace(".doc", "").replace(".docx", "").replace(".wps", "")
            out_name = f"{path_prefix}.txt"
            out_path = OUT_DIR / out_name
            out_path.write_text(text, encoding="utf-8")

            lines_count = len([l for l in text.split("\n") if l.strip()])
            char_count = len(text)
            print(f"✅ {char_count}字符/{lines_count}行")

            results.append({
                "file": rel,
                "ext": ext,
                "chars": char_count,
                "lines": lines_count,
                "saved_to": str(out_path),
            })
        except Exception as e:
            print(f"❌ {e}")
            results.append({"file": rel, "ext": ext, "error": str(e)})

    # 汇总
    print("\n" + "=" * 60)
    print("提取汇总")
    print("=" * 60)
    ok = [r for r in results if "error" not in r]
    fail = [r for r in results if "error" in r]
    total_chars = sum(r.get("chars", 0) for r in ok)
    print(f"成功: {len(ok)} 个, 失败: {len(fail)} 个, 总字符: {total_chars}")

    if fail:
        print("\n失败文件:")
        for r in fail:
            print(f"  ❌ {r['file']}: {r.get('error', '')}")


if __name__ == "__main__":
    main()