#!/usr/bin/env python3
"""
Excel 文档内容提取与汇总分析
"""
import os, sys
from pathlib import Path
from datetime import datetime

PDFS_DIR = Path("/home/scada/SmartTwinRes-skills/pdfs")
OUT_DIR = Path("/home/scada/SmartTwinRes-skills/pdf_text_analysis/excel_extracted")
OUT_DIR.mkdir(parents=True, exist_ok=True)

def extract_xlsx(path: Path) -> str:
    """提取 xlsx 内容"""
    import openpyxl
    lines = []
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        lines.append(f"===== Sheet: {sheet} ({ws.max_row}行 x {ws.max_column}列) =====")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def extract_xls(path: Path) -> str:
    """提取旧版 xls 内容（calamine 引擎，容错性更好）"""
    from python_calamine import CalamineWorkbook
    lines = []
    wb = CalamineWorkbook.from_path(str(path))
    for sn in wb.sheet_names:
        rows = wb.get_sheet_by_name(sn).to_python()
        lines.append(f"===== Sheet: {sn} ({len(rows)}行) =====")
        for row in rows:
            cells = [str(c).strip() if c is not None else "" for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def main():
    print("=" * 60)
    print("Excel 文档提取")
    print("=" * 60)

    xls_files = sorted(PDFS_DIR.rglob("*.xls"))
    xlsx_files = sorted(PDFS_DIR.rglob("*.xlsx"))
    all_files = sorted(xls_files + xlsx_files, key=lambda p: str(p))
    print(f"共 {len(all_files)} 个 Excel 文件\n")

    results = []
    for fp in all_files:
        rel = str(fp.relative_to(PDFS_DIR))
        ext = fp.suffix.lower()
        print(f"[{all_files.index(fp)+1}/{len(all_files)}] {rel}", end=" ", flush=True)

        try:
            if ext == ".xlsx":
                text = extract_xlsx(fp)
            elif ext == ".xls":
                text = extract_xls(fp)
            else:
                continue

            # 保存（文件名带路径前缀，避免不同目录同名文件互相覆盖）
            path_prefix = str(fp.relative_to(PDFS_DIR)).replace("/", "_").replace(".xls", "").replace(".xlsx", "")
            out_name = f"{path_prefix}.txt"
            out_path = OUT_DIR / out_name
            out_path.write_text(text, encoding="utf-8")

            # 统计
            lines_count = len([l for l in text.split("\n") if l.strip()])
            char_count = len(text)
            print(f"✅ {char_count}字符/{lines_count}行")

            results.append({
                "file": rel,
                "ext": ext,
                "chars": char_count,
                "lines": lines_count,
                "saved_to": str(out_path),
            })
        except Exception as e:
            print(f"❌ {e}")
            results.append({"file": rel, "ext": ext, "error": str(e)})

    # 汇总
    print("\n" + "=" * 60)
    print("提取汇总")
    print("=" * 60)
    ok = [r for r in results if "error" not in r]
    fail = [r for r in results if "error" in r]
    total_chars = sum(r.get("chars", 0) for r in ok)
    print(f"成功: {len(ok)} 个, 失败: {len(fail)} 个, 总字符: {total_chars}")

    if fail:
        print("\n失败文件:")
        for r in fail:
            print(f"  ❌ {r['file']}: {r['error']}")


if __name__ == "__main__":
    main()